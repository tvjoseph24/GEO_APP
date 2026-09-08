"""Ingestion tests: doc parsing (free, self-generated fixtures) + web (network)."""
import os, sys, tempfile
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from ingestion import ingest_doc, ingest_url, ingest_folder, RawSource

def test_txt_ingestion():
    with tempfile.NamedTemporaryFile("w", suffix=".txt", delete=False) as f:
        f.write("Aveeno uses colloidal oatmeal. Fragrance-free.")
        path = f.name
    src = ingest_doc(path)
    assert src.source_type == "doc"
    assert "colloidal oatmeal" in src.text
    os.unlink(path)
    print("✅ txt ingestion")

def test_docx_ingestion():
    try:
        import docx
    except ImportError:
        print("⚠️  python-docx not installed — skip docx test"); return
    d = docx.Document()
    d.add_paragraph("NEA Seal of Acceptance.")
    d.add_paragraph("Steroid-free formula.")
    path = tempfile.mktemp(suffix=".docx"); d.save(path)
    src = ingest_doc(path)
    assert "NEA" in src.text and "Steroid-free" in src.text
    os.unlink(path)
    print("✅ docx ingestion")

def test_bytes_ingestion():
    src = ingest_doc(b"1% colloidal oatmeal", name="spec.txt")
    assert "colloidal oatmeal" in src.text
    print("✅ bytes ingestion (upload path)")

def test_folder_ingestion():
    d = tempfile.mkdtemp()
    for i, txt in enumerate(["claim one", "claim two"]):
        open(os.path.join(d, f"f{i}.txt"), "w").write(txt)
    srcs = ingest_folder(d)
    assert len(srcs) == 2
    print("✅ folder ingestion")

def test_web_ingestion_network():
    """Requires network. Uses a stable page."""
    src = ingest_url("https://example.com", timeout=10)
    if src.meta.get("error"):
        print(f"⚠️  web test skipped (network): {src.meta['error'][:60]}")
        return
    assert src.source_type == "web" and src.meta.get("status") == 200
    assert len(src.text) > 0
    print(f"✅ web ingestion (domain={src.meta.get('domain')}, chars={src.meta.get('chars')})")

def test_web_bad_url_graceful():
    src = ingest_url("https://this-domain-does-not-exist-999.xyz", timeout=5)
    assert "error" in src.meta   # should fail gracefully, not crash
    print("✅ web bad-url handled gracefully")

if __name__ == "__main__":
    print("\n── INGESTION TESTS ──")
    for fn in [test_txt_ingestion, test_docx_ingestion, test_bytes_ingestion,
               test_folder_ingestion, test_web_ingestion_network,
               test_web_bad_url_graceful]:
        try: fn()
        except Exception as e: print(f"❌ {fn.__name__}: {type(e).__name__}: {e}")