"""Word renderer. House styles generated programmatically — no template required."""
from __future__ import annotations

import io

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from report.model import ReportModel
from theme import PRIMARY, INK as _INK, INK_MUTED as _MUTED

def _rgb(h):
    h = h.lstrip("#")
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))

INK   = _rgb(_INK)
BRAND = _rgb(PRIMARY)
MUTED = _rgb(_MUTED)

TONE_FILL = {"info": "EEF3FF", "warn": "FFF6E5", "danger": "FDECEC", "good": "EAF7EF"}
TONE_EDGE = {"info": "0B5FFF", "warn": "E8A33D", "danger": "D64545", "good": "1B9E5A"}

def _shade(cell, hexcolor):
    el = OxmlElement("w:shd")
    el.set(qn("w:val"), "clear")
    el.set(qn("w:fill"), hexcolor)
    cell._tc.get_or_add_tcPr().append(el)

def _borders(cell, colour, width=8, sides=("left",)):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for side in sides:
        e = OxmlElement(f"w:{side}")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), str(width))
        e.set(qn("w:color"), colour)
        borders.append(e)
    tcPr.append(borders)

def _styles(doc):
    n = doc.styles["Normal"]
    n.font.name = "Calibri"
    n.font.size = Pt(10)
    n.font.color.rgb = INK
    n.paragraph_format.space_after = Pt(6)
    n.paragraph_format.line_spacing = 1.12
    for i, (sz, col, before) in enumerate(
            [(20, BRAND, 20), (14, INK, 16), (11.5, INK, 12)], start=1):
        st = doc.styles[f"Heading {i}"]
        st.font.name = "Calibri"
        st.font.size = Pt(sz)
        st.font.bold = True
        st.font.color.rgb = col
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(6)
    for name, sz, col, italic in (("GEO Caption", 8.5, MUTED, True),
                                  ("GEO Small", 9, MUTED, False),
                                  ("GEO Mono", 8.5, INK, False)):
        try:
            st = doc.styles.add_style(name, 1)
        except Exception:
            st = doc.styles[name]
        st.font.name = "Consolas" if name == "GEO Mono" else "Calibri"
        st.font.size = Pt(sz)
        st.font.italic = italic
        st.font.color.rgb = col
        st.paragraph_format.space_after = Pt(8)

def _md(par, text):
    """Minimal **bold** / _italic_ inline handling."""
    import re
    for tok in re.split(r"(\*\*[^*]+\*\*|_[^_]+_)", str(text or "")):
        if not tok:
            continue
        if tok.startswith("**") and tok.endswith("**"):
            par.add_run(tok[2:-2]).bold = True
        elif tok.startswith("_") and tok.endswith("_") and len(tok) > 2:
            par.add_run(tok[1:-1]).italic = True
        else:
            par.add_run(tok)

def _callout(doc, text, tone):
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = t.cell(0, 0)
    _shade(c, TONE_FILL.get(tone, "EEF3FF"))
    _borders(c, TONE_EDGE.get(tone, "0B5FFF"), 18, ("left",))
    c.paragraphs[0].style = doc.styles["Normal"]
    _md(c.paragraphs[0], text)
    doc.add_paragraph()

def _table(doc, rows, caption=""):
    rows = [r for r in (rows or []) if isinstance(r, dict)]
    if not rows:
        return
    cols = list(rows[0].keys())
    t = doc.add_table(rows=1, cols=len(cols))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j, c in enumerate(cols):
        cell = t.rows[0].cells[j]
        cell.text = ""
        r = cell.paragraphs[0].add_run(str(c))
        r.bold = True
        r.font.size = Pt(8.5)
        _shade(cell, "F2F4F7")
    for row in rows:
        cells = t.add_row().cells
        for j, c in enumerate(cols):
            v = row.get(c)
            cells[j].text = ""
            run = cells[j].paragraphs[0].add_run(
                "" if v is None else str(v)[:600])
            run.font.size = Pt(8.5)
    if caption:
        p = doc.add_paragraph(caption, style="GEO Caption")
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    doc.add_paragraph()

def _kpi_row(doc, kpis):
    t = doc.add_table(rows=2, cols=len(kpis))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j, k in enumerate(kpis):
        top = t.rows[0].cells[j]
        _shade(top, "F2F4F7")
        top.text = ""
        rr = top.paragraphs[0].add_run(str(k.get("label")))
        rr.bold = True
        rr.font.size = Pt(9)
        p2 = top.add_paragraph()
        r = p2.add_run(str(k.get("value")))
        r.bold = True
        r.font.size = Pt(15)
        r.font.color.rgb = BRAND
        bot = t.rows[1].cells[j]
        bot.text = ""
        n = bot.paragraphs[0].add_run(str(k.get("note", "")))
        n.font.size = Pt(8)
        n.font.color.rgb = MUTED
        if k.get("fact"):
            fr = bot.paragraphs[0].add_run(f"  [{k['fact']}]")
            fr.font.size = Pt(7)
            fr.font.color.rgb = MUTED
    doc.add_paragraph()

def _cover(doc, meta, coverage, provisional, reasons):
    for _ in range(4):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("GEO INTELLIGENCE REPORT")
    r.bold = True
    r.font.size = Pt(11)
    r.font.color.rgb = MUTED
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(str(meta.get("brand", "")).title())
    r.bold = True
    r.font.size = Pt(34)
    r.font.color.rgb = BRAND
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(f"{meta.get('parent', '').title()} · {meta.get('category', '')}")
    r.font.size = Pt(13)
    r.font.color.rgb = MUTED
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(f"How generative AI engines describe and recommend the brand\n"
                  f"{', '.join(coverage.get('engines') or [])} · "
                  f"{coverage.get('n_runs', 0):,} measured runs")
    r.font.size = Pt(9.5)
    r.font.color.rgb = MUTED
    doc.add_paragraph()
    if provisional:
        _callout(doc, "**PROVISIONAL ANALYSIS** — " + " ".join(reasons), "danger")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(f"Generated {meta.get('generated_utc', '')[:10]} · "
                  f"profile: {meta.get('profile')} · "
                  f"v{meta.get('generator_version')}")
    r.font.size = Pt(8)
    r.font.color.rgb = MUTED
    doc.add_page_break()

def _footer(doc, meta):
    for sec in doc.sections:
        p = sec.footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(f"GEO Intelligence Report · {meta.get('brand', '').title()} · "
                      f"{meta.get('generated_utc', '')[:10]} · "
                      f"point-in-time measurement")
        r.font.size = Pt(7.5)
        r.font.color.rgb = MUTED

def render_docx(rm: ReportModel, audience="both", profile=None) -> bytes:
    profile = profile or rm.meta.get("profile", "internal")
    doc = Document()
    _styles(doc)
    for sec in doc.sections:
        sec.left_margin = sec.right_margin = Inches(0.9)
        sec.top_margin = sec.bottom_margin = Inches(0.85)

    _cover(doc, rm.meta, rm.coverage, rm.provisional, rm.provisional_reasons)

    sections = rm.filtered(audience, profile)
    doc.add_heading("Contents", 1)
    for i, s in enumerate(sections, 1):
        p = doc.add_paragraph(style="GEO Small")
        p.add_run(f"{i}.  {s.title}")
    doc.add_page_break()

    for i, s in enumerate(sections, 1):
        doc.add_heading(f"{i}. {s.title}", 1)
        if s.subtitle:
            doc.add_paragraph(s.subtitle, style="GEO Small")
        for b in s.blocks:
            if b.kind == "heading":
                doc.add_heading(str(b.payload), min(b.level + 1, 3))
            elif b.kind == "prose":
                par = doc.add_paragraph()
                _md(par, b.payload)
                if b.edited and profile == "internal":
                    mark = par.add_run("  [edited]")
                    mark.font.size = Pt(7)
                    mark.font.color.rgb = MUTED
            elif b.kind == "bullets":
                for it in (b.payload or []):
                    _md(doc.add_paragraph(style="List Bullet"), it)
            elif b.kind == "table":
                _table(doc, b.payload, b.caption)
            elif b.kind == "coverage":
                _table(doc, [{"Parameter": k, "Value": v}
                             for k, v in (b.payload or {}).items()], b.caption)
            elif b.kind == "evidence_key":
                _table(doc, b.payload, "How to read the evidence classes used throughout.")
            elif b.kind == "kpi_row":
                _kpi_row(doc, b.payload)
            elif b.kind == "callout":
                _callout(doc, b.payload, b.tone)
            elif b.kind == "chart" and b.payload:
                doc.add_picture(io.BytesIO(b.payload), width=Inches(6.3))
                if b.caption:
                    doc.add_paragraph(b.caption, style="GEO Caption")
            elif b.kind == "quote":
                t = doc.add_table(rows=1, cols=1)
                c = t.cell(0, 0)
                _shade(c, "F7F8FA")
                _borders(c, "8C9196", 12, ("left",))
                c.text = ""
                r = c.paragraphs[0].add_run(str(b.payload))
                r.font.size = Pt(8.5)
                r.font.name = "Consolas"
                if b.caption:
                    doc.add_paragraph(b.caption, style="GEO Caption")
                doc.add_paragraph()
            elif b.kind == "pagebreak":
                doc.add_page_break()
        if i < len(sections):
            doc.add_page_break()

    _footer(doc, rm.meta)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()