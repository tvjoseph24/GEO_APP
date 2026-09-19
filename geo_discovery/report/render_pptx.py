"""Executive deck — a FILTER over the report IR, not a compressed document.

Narrative arc: what did we find → why it matters → what we proved → what we ruled out →
what to do → how we will know. Charts reuse the same PNGs as the Word report.
"""
from __future__ import annotations

import io

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Inches, Pt
from theme import PRIMARY, INK as _INK, INK_MUTED as _MUTED, DANGER, GOOD as _GOOD


def _rgb(h):
    h = h.lstrip("#")
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))

BRAND = _rgb(PRIMARY)
INK   = _rgb(_INK)
MUTED = _rgb(_MUTED)
BAD   = _rgb(DANGER)
GOOD  = _rgb(_GOOD)

W, H = Inches(13.333), Inches(7.5)

def _deck():
    p = Presentation()
    p.slide_width, p.slide_height = W, H
    return p

def _blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])

def _tx(slide, text, l, t, w, h, size=18, bold=False, colour=INK,
        align=PP_ALIGN.LEFT, wrap=True):
    box = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = wrap
    lines = str(text or "").split("\n")
    for i, line in enumerate(lines):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = align
        r = para.add_run()
        r.text = line
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = colour
        r.font.name = "Calibri"
    return box

def _title(slide, text, kicker=""):
    if kicker:
        _tx(slide, kicker.upper(), 0.7, 0.42, 11, 0.3, 11, True, MUTED)
    _tx(slide, text, 0.7, 0.72, 12, 0.9, 30, True, INK)
    ln = slide.shapes.add_shape(1, Inches(0.7), Inches(1.62), Inches(1.6), Emu(22860))
    ln.fill.solid()
    ln.fill.fore_color.rgb = BRAND
    ln.line.fill.background()

def _stat(slide, l, t, value, label, note="", colour=BRAND, w=2.9):
    _tx(slide, value, l, t, w, 0.75, 40, True, colour)
    _tx(slide, label, l, t + 0.78, w, 0.32, 13, True, INK)
    if note:
        _tx(slide, note, l, t + 1.10, w, 0.9, 9.5, False, MUTED)

def _bullets(slide, items, l, t, w, size=15, gap=0.52, colour=INK, limit=7):
    for i, it in enumerate(items[:limit]):
        _tx(slide, f"•  {it}", l, t + i * gap, w, gap, size, False, colour)

def _chart(slide, png, l, t, w):
    if png:
        slide.shapes.add_picture(io.BytesIO(png), Inches(l), Inches(t), width=Inches(w))

def _find(rm, section_id, kind, index=0):
    """Pull a specific block out of the IR — the deck reuses, never recomputes."""
    for s in rm.sections:
        if s.id != section_id:
            continue
        hits = [b for b in s.blocks if b.kind == kind]
        if len(hits) > index:
            return hits[index]
    return None

def render_pptx(rm, profile=None) -> bytes:
    profile = profile or rm.meta.get("profile", "internal")
    prs = _deck()
    cov, meta = rm.coverage, rm.meta
    ev = {}
    for s in rm.sections:
        if s.id == "exec":
            for b in s.blocks:
                if b.kind == "kpi_row":
                    ev = {k["label"]: k for k in b.payload}

    # ---- 1 · cover ----
    sl = _blank(prs)
    _tx(sl, "GEO INTELLIGENCE", 0.7, 2.3, 11, 0.4, 13, True, MUTED)
    _tx(sl, str(meta.get("brand", "")).title(), 0.7, 2.7, 11, 1.3, 52, True, BRAND)
    _tx(sl, f"{meta.get('parent','').title()} · {meta.get('category','')}",
        0.7, 4.0, 11, 0.5, 20, False, MUTED)
    _tx(sl, f"How generative AI engines describe and recommend the brand\n"
            f"{', '.join(cov.get('engines') or [])} · "
            f"{cov.get('n_runs', 0):,} measured runs · "
            f"{meta.get('generated_utc','')[:10]}",
        0.7, 4.7, 11, 1.1, 12, False, MUTED)
    if rm.provisional:
        _tx(sl, "PROVISIONAL ANALYSIS", 0.7, 6.3, 11, 0.4, 14, True, BAD)

    # ---- 2 · the situation ----
    sl = _blank(prs)
    _title(sl, "The situation in machine reality", "Diagnosis")
    order = ["Unprompted visibility", "Retrieval gap", "Citation control", "Voice control"]
    cols = [(0.7, BAD), (3.9, INK), (7.1, BAD), (10.3, BAD)]
    for name, (x, c) in zip(order, cols):
        k = ev.get(name)
        if k:
            _stat(sl, x, 2.1, k["value"], k["label"], k.get("note", ""), c)
    ovs = _find(rm, "exec", "callout")
    if ovs:
        _tx(sl, str(ovs.payload)[:420], 0.7, 5.3, 12, 1.6, 14, False, INK)

    # ---- 3 · retrieval gap chart ----
    ch = _find(rm, "exec", "chart")
    if ch:
        sl = _blank(prs)
        _title(sl, "The brand wins when present — and is absent", "The core finding")
        _chart(sl, ch.payload, 2.4, 1.9, 8.5)
        _tx(sl, ch.caption[:340], 0.7, 6.2, 12, 1.0, 12, False, MUTED)

    # ---- 4 · principal findings ----
    sl = _blank(prs)
    _title(sl, "What we found", "Evidence")
    bl = _find(rm, "exec", "bullets")
    if bl:
        _bullets(sl, [str(x)[:190] for x in bl.payload], 0.7, 2.0, 12, 14, 0.62, limit=6)

    # ---- 5 · competitive position ----
    rich = _find(rm, "delta", "chart", 0)
    if rich:
        sl = _blank(prs)
        _title(sl, "Competitors argue harder in less space", "Competitive position")
        _chart(sl, rich.payload, 2.4, 1.9, 8.5)
        _tx(sl, rich.caption[:320], 0.7, 6.2, 12, 1.0, 12, False, MUTED)

    # ---- 6 · what we proved ----
    forest = _find(rm, "experiment", "chart")
    if forest:
        sl = _blank(prs)
        _title(sl, "What we proved — and what we did not", "Experiment")
        _chart(sl, forest.payload, 2.4, 1.85, 8.5)
        _tx(sl, "Bars at zero are well-powered nulls: tested, no effect. These are "
                "results, not missing data.", 0.7, 6.3, 12, 0.8, 12, False, MUTED)

    # ---- 7 · what we rule out ----
    ruled = None
    for s in rm.sections:
        if s.id == "exec":
            tabs = [b for b in s.blocks if b.kind == "table"]
            if len(tabs) >= 2:
                ruled = tabs[1]
    if ruled and ruled.payload:
        sl = _blank(prs)
        _title(sl, "What we will not spend on", "Evidence-based exclusions")
        for i, r in enumerate(ruled.payload[:3]):
            y = 2.0 + i * 1.6
            _tx(sl, f"✕  {r.get('Excluded','')}"[:96], 0.7, y, 12, 0.4, 16, True, BAD)
            _tx(sl, str(r.get("Reason", ""))[:300], 1.1, y + 0.44, 11.4, 1.0, 12,
                False, INK)
            _tx(sl, str(r.get("Basis", "")), 1.1, y + 1.22, 11.4, 0.3, 10, False, MUTED)

    # ---- 8 · big bets ----
    strat = rm.meta.get("_big_bets") or []
    if not strat:
        for s in rm.sections:
            if s.id == "exec":
                for b in s.blocks:
                    if b.kind == "table" and b.payload and "Priority" in b.payload[0]:
                        strat = b.payload
                        break
    if strat:
        sl = _blank(prs)
        _title(sl, "Where to act", "Priorities")
        for i, b in enumerate(strat[:3]):
            y = 1.95 + i * 1.72
            _tx(sl, f"{i+1}", 0.7, y, 0.5, 0.6, 30, True, BRAND)
            _tx(sl, str(b.get("Priority", ""))[:88], 1.3, y, 8.3, 0.42, 17, True, INK)
            _tx(sl, str(b.get("Evidence", "")), 9.9, y + 0.04, 3.0, 0.32, 11, True,
                GOOD if "Proven" in str(b.get("Evidence", "")) else MUTED)
            _tx(sl, str(b.get("Expected effect", ""))[:230], 1.3, y + 0.48, 11.3, 0.7,
                11.5, False, INK)
            _tx(sl, str(b.get("Why now", ""))[:210], 1.3, y + 1.14, 11.3, 0.5, 10,
                False, MUTED)

    # ---- 9 · substantiation exposure ----
    sub = None
    for s in rm.sections:
        if s.id == "exec":
            calls = [b for b in s.blocks if b.kind == "callout" and b.tone == "danger"]
            if calls:
                sub = calls[-1]
    if sub:
        sl = _blank(prs)
        _title(sl, "Claims the engines make that we have not substantiated",
               "Compliance exposure")
        _tx(sl, str(sub.payload).replace("**", "")[:640], 0.7, 2.1, 12, 3.0, 16,
            False, INK)
        _tx(sl, "Lowest-cost, highest-value work available — and in a regulated "
                "category, a compliance matter.", 0.7, 5.6, 12, 0.8, 14, True, BAD)

    # ---- 10 · roadmap ----
    waves = None
    for s in rm.sections:
        if s.id == "roadmap":
            for b in s.blocks:
                if b.kind == "table" and b.payload and "Ref" in (b.payload[0] or {}):
                    waves = waves or []
                    waves.append((b.caption, b.payload))
    if waves:
        sl = _blank(prs)
        _title(sl, "Sequenced roadmap", "Delivery")
        for i, (cap, rows) in enumerate(waves[:4]):
            x = 0.7 + i * 3.15
            _tx(sl, str(cap).split("·")[0][:26], x, 1.95, 3.0, 0.4, 14, True, BRAND)
            _bullets(sl, [f"{r.get('Ref')} {str(r.get('Action',''))[:34]}"
                          for r in rows[:5]], x, 2.45, 3.0, 10, 0.42)

    # ---- 11 · projection ----
    proj = None
    for s in rm.sections:
        if s.id == "roadmap":
            chs = [b for b in s.blocks if b.kind == "chart"]
            if chs:
                proj = chs[-1]
    if proj:
        sl = _blank(prs)
        _title(sl, "Projected trajectory", "Illustrative, not a forecast")
        _chart(sl, proj.payload, 2.4, 1.9, 8.5)
        _tx(sl, "Contributions are discounted by evidence class and by overlap within a "
                "mechanism. Actions with no measured basis contribute nothing.",
            0.7, 6.2, 12, 1.0, 12, False, MUTED)

    # ---- 12 · how we will know ----
    sl = _blank(prs)
    _title(sl, "How we will know it worked", "Measurement")
    mc = []
    for s in rm.sections:
        if s.id == "roadmap":
            for b in s.blocks:
                if b.kind == "table" and b.payload and \
                        (b.payload[0] or {}).get("Field") == "Queries to re-audit":
                    mc = b.payload
                    break
            if mc:
                break
    if mc:
        _bullets(sl, [f"{r.get('Field')}: {r.get('Detail')}" for r in mc], 0.7, 2.0, 12,
                 14, 0.55, limit=8)
    _tx(sl, "Every wave carries a measurement contract: which queries to re-audit, at "
            "what sample size, on which engines, by when, and what result would stop the "
            "work.", 0.7, 6.2, 12, 1.0, 13, False, MUTED)

    # ---- 13 · scope and limits ----
    sl = _blank(prs)
    _title(sl, "Scope and limitations", "Read this before acting")
    lims = [
        f"Engines tested: {', '.join(cov.get('engines') or [])}. "
        f"{cov.get('n_runs', 0):,} runs.",
        f"Query categories audited: {len(cov.get('intents_audited') or [])} of "
        f"{len(cov.get('intents_planned') or [])}. Not audited: "
        f"{', '.join(cov.get('intents_not_audited') or ['none'])}.",
        "Visual, multimodal and conversational behaviour was not assessed."
        if cov.get("vision_verified") is not True else
        "Visual extraction was verified.",
        "Only retrieval mode was audited — memory-driven and retrieval-driven behaviour "
        "cannot be separated." if not cov.get("memory_vs_retrieval_separated") else
        "Memory and retrieval modes were audited separately.",
        "Experiments inject content into context, so they measure how content performs "
        "once found — not whether it is found.",
        "AI engines and their indexes change continuously. This is a point-in-time "
        "measurement.",
    ]
    _bullets(sl, lims, 0.7, 2.0, 12, 13.5, 0.72, MUTED, limit=6)
    _tx(sl, f"{'APPROVED' if meta.get('signoff') == 'approved' else 'DRAFT'} · "
            f"profile: {profile} · generated {meta.get('generated_utc','')[:10]}",
        0.7, 6.8, 12, 0.4, 10, True, MUTED)

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()